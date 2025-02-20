from docx import Document

# Create a new Word document
doc = Document()

# Add the content to the document
doc.add_heading('SAURABH RAJENDRA KHATAVKAR', level=1)

# Personal Details
doc.add_heading('Personal Details', level=2)
doc.add_paragraph("Marital Status: Never Married")
doc.add_paragraph("Birth Date: 28 March 1992")
doc.add_paragraph("Birth Time: 08:47 AM")
doc.add_paragraph("Birth Place: Gulbarga, Karnataka")
doc.add_paragraph("Rashi: Makar")
doc.add_paragraph("Nakshatra: Uttarashadha")
doc.add_paragraph("Charan: 3")
doc.add_paragraph("Height: 5'7\"")
doc.add_paragraph("Blood Group: O+ve")

# Education Details
doc.add_heading('Education Details', level=2)
doc.add_paragraph("Degree: MTech (Computer Science)")
doc.add_paragraph("College: IIT Guwahati")

# Occupation Details
doc.add_heading('Occupation Details', level=2)
doc.add_paragraph("Designation: Software Engineer III at MNC, Bangalore")
doc.add_paragraph("Income: 45.5 LPA")

# Family Details
doc.add_heading('Family Details', level=2)
doc.add_paragraph("Father: Dr. Rajendra Hari Khatavkar")
doc.add_paragraph("Mother: Dr. Rajaswini Rajendra Khatavkar")
doc.add_paragraph("Mamekul: Mahindrakar, Gulbarga")
doc.add_paragraph("Siblings: 1 Brother, BE & MBA, Married")

# Contact Details
doc.add_heading('Contact Details', level=2)
doc.add_paragraph("Address: Airoli, Navi Mumbai, 400708")
doc.add_paragraph("Mobile: 9223211959, 9223522959")

# Expectations
doc.add_heading('Expectations', level=2)
doc.add_paragraph("Understanding, Well-Educated, Preferably Working in IT Sector & Engineer/MBA")

# Save the document
output_path = "Z:\pythonfile/Saurabh_Rajendra_Khatavkar_Profile.docx"
doc.save(output_path)

print(f"Document saved at: {output_path}")
